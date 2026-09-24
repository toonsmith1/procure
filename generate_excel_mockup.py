import json
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load mockup JSON data
with open("mockup-data.json", "r", encoding="utf-8") as f:
    data = json.load(f)

contracts = data["contracts"]
reports = data["reports"]
report_items = data["reportItems"]

wb = Workbook()

# Define styles
header_font = Font(name="TH Sarabun New", size=14, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

report_title_font = Font(name="TH Sarabun New", size=18, bold=True, color="0F172A")
report_sub_font = Font(name="TH Sarabun New", size=14, bold=True, color="334155")
report_header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
report_header_font = Font(name="TH Sarabun New", size=13, bold=True, color="FFFFFF")

data_font = Font(name="TH Sarabun New", size=13)
data_font_bold = Font(name="TH Sarabun New", size=13, bold=True)

thin_border_side = Side(border_style="thin", color="CBD5E1")
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

money_format = '#,##0.00'

# -----------------------------------------------------------------------------
# 1. Sheet: Contracts (Matches Google Sheet Schema)
# -----------------------------------------------------------------------------
ws_contracts = wb.active
ws_contracts.title = "Contracts"
ws_contracts.freeze_panes = "A2"

contract_headers = [
    "id", "type", "document_no", "egp_project_no", "control_no", "gfmis_po",
    "document_date", "vendor_name", "tax_id", "project_name", "budget",
    "median_price", "offered_price", "agreed_price", "contract_amount",
    "start_date", "end_date", "guarantee_type", "guarantee_amount",
    "status", "note", "created_at", "updated_at"
]

ws_contracts.append(contract_headers)
for col_idx in range(1, len(contract_headers) + 1):
    cell = ws_contracts.cell(row=1, column=col_idx)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

numeric_cols = {"budget", "median_price", "offered_price", "agreed_price", "contract_amount", "guarantee_amount"}

for row_idx, item in enumerate(contracts, start=2):
    for col_idx, col_name in enumerate(contract_headers, start=1):
        val = item.get(col_name, "")
        cell = ws_contracts.cell(row=row_idx, column=col_idx)
        cell.font = data_font
        cell.border = thin_border
        
        if col_name in numeric_cols:
            try:
                cell.value = float(val) if val else 0.0
                cell.number_format = money_format
                cell.alignment = Alignment(horizontal="right", vertical="center")
            except ValueError:
                cell.value = val
                cell.alignment = Alignment(horizontal="left", vertical="center")
        elif col_name in ["id", "document_no", "egp_project_no", "control_no", "gfmis_po", "tax_id", "document_date", "start_date", "end_date", "status"]:
            cell.value = str(val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.value = str(val)
            cell.alignment = Alignment(horizontal="left", vertical="center")

# Auto-adjust column widths
for col in ws_contracts.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or "")
        max_len = max(max_len, len(val_str.encode("utf-8", "ignore")) // 2)
    ws_contracts.column_dimensions[col_letter].width = max(max_len + 4, 12)

# -----------------------------------------------------------------------------
# 2. Sheet: MonthlyReports (Matches Google Sheet Schema)
# -----------------------------------------------------------------------------
ws_reports = wb.create_sheet(title="MonthlyReports")
ws_reports.freeze_panes = "A2"

report_headers = [
    "id", "year", "month", "report_name", "published_date", "published_url",
    "status", "note", "created_at", "updated_at"
]

ws_reports.append(report_headers)
for col_idx in range(1, len(report_headers) + 1):
    cell = ws_reports.cell(row=1, column=col_idx)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

for row_idx, item in enumerate(reports, start=2):
    for col_idx, col_name in enumerate(report_headers, start=1):
        val = item.get(col_name, "")
        cell = ws_reports.cell(row=row_idx, column=col_idx)
        cell.value = str(val)
        cell.font = data_font
        cell.border = thin_border
        if col_name in ["id", "year", "month", "published_date", "status"]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

for col in ws_reports.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or "")
        max_len = max(max_len, len(val_str.encode("utf-8", "ignore")) // 2)
    ws_reports.column_dimensions[col_letter].width = max(max_len + 4, 14)

# -----------------------------------------------------------------------------
# 3. Sheet: MonthlyReportItems (Matches Google Sheet Schema)
# -----------------------------------------------------------------------------
ws_items = wb.create_sheet(title="MonthlyReportItems")
ws_items.freeze_panes = "A2"

item_headers = ["id", "report_id", "contract_id", "created_at"]
ws_items.append(item_headers)
for col_idx in range(1, len(item_headers) + 1):
    cell = ws_items.cell(row=1, column=col_idx)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

for row_idx, item in enumerate(report_items, start=2):
    for col_idx, col_name in enumerate(item_headers, start=1):
        val = item.get(col_name, "")
        cell = ws_items.cell(row=row_idx, column=col_idx)
        cell.value = str(val)
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

for col in ws_items.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or "")
        max_len = max(max_len, len(val_str))
    ws_items.column_dimensions[col_letter].width = max(max_len + 6, 16)

# -----------------------------------------------------------------------------
# 4. Sheet: สขร.1_มกราคม_2569 (Official Thai Government Procurement Format)
# -----------------------------------------------------------------------------
ws_skr1 = wb.create_sheet(title="สขร.1_มกราคม_2569")

# Header titles
ws_skr1.merge_cells("A1:I1")
ws_skr1["A1"] = "แบบสรุปผลการดำเนินการจัดซื้อจัดจ้างในรอบเดือน มกราคม 2569"
ws_skr1["A1"].font = report_title_font
ws_skr1["A1"].alignment = Alignment(horizontal="center", vertical="center")

ws_skr1.merge_cells("A2:I2")
ws_skr1["A2"] = "ศูนย์ข้อมูลเกษตรแห่งชาติ สำนักงานเศรษฐกิจการเกษตร"
ws_skr1["A2"].font = report_sub_font
ws_skr1["A2"].alignment = Alignment(horizontal="center", vertical="center")

ws_skr1.row_dimensions[1].height = 28
ws_skr1.row_dimensions[2].height = 22
ws_skr1.row_dimensions[4].height = 36

skr1_cols = [
    "ลำดับ",
    "งานที่จัดซื้อหรือจัดจ้าง",
    "วงเงินที่จัดซื้อ\nหรือจัดจ้าง (บาท)",
    "ราคากลาง\n(บาท)",
    "วิธีซื้อหรือจ้าง",
    "รายชื่อผู้เสนอราคา\nและราคาที่เสนอ",
    "ผู้ได้รับการคัดเลือกและราคา\nที่ตกลงซื้อหรือจ้าง",
    "เหตุผลที่คัดเลือก\nโดยสรุป",
    "เลขที่และวันที่ของสัญญา\nหรือข้อตกลงในการซื้อหรือจ้าง"
]

for col_idx, col_title in enumerate(skr1_cols, start=1):
    cell = ws_skr1.cell(row=4, column=col_idx)
    cell.value = col_title
    cell.font = report_header_font
    cell.fill = report_header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

# Filter items for Jan 2569 (skr1-2569-01)
jan_contract_ids = [it["contract_id"] for it in report_items if it["report_id"] == "skr1-2569-01"]
jan_contracts = [c for c in contracts if c["id"] in jan_contract_ids]

start_data_row = 5
total_budget = 0.0
total_agreed = 0.0

for idx, c in enumerate(jan_contracts, start=1):
    cur_row = start_data_row + idx - 1
    ws_skr1.row_dimensions[cur_row].height = 30
    
    b_val = float(c.get("budget", 0))
    m_val = float(c.get("median_price", 0))
    off_val = float(c.get("offered_price", 0))
    agr_val = float(c.get("agreed_price", 0))
    
    total_budget += b_val
    total_agreed += agr_val
    
    method = "เฉพาะเจาะจง" if b_val <= 500000 else "e-Bidding (ประกวดราคาอิเล็กทรอนิกส์)"
    doc_no = c.get("document_no", "")
    doc_date = c.get("document_date", "")
    d_parts = doc_date.split("-")
    thai_date_str = f"{d_parts[2]}/{d_parts[1]}/{int(d_parts[0])+543}" if len(d_parts) == 3 else doc_date
    
    row_vals = [
        idx,
        c.get("project_name", ""),
        b_val,
        m_val,
        method,
        f"{c.get('vendor_name', '')}\nเสนอราคา {off_val:,.2f} บาท",
        f"{c.get('vendor_name', '')}\nตกลงราคา {agr_val:,.2f} บาท",
        "เป็นผู้มีคุณสมบัติถูกต้องตามเงื่อนไขและเสนอราคาเหมาะสม",
        f"เลขที่ {doc_no}\nลงวันที่ {thai_date_str}"
    ]
    
    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws_skr1.cell(row=cur_row, column=col_idx)
        cell.value = val
        cell.font = data_font
        cell.border = thin_border
        
        if col_idx == 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [3, 4]:
            cell.number_format = money_format
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx in [5, 8]:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Summary Total Row
summary_row = start_data_row + len(jan_contracts)
ws_skr1.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=2)
sum_label = ws_skr1.cell(row=summary_row, column=1)
sum_label.value = f"รวมทั้งสิ้น ({len(jan_contracts)} รายการ)"
sum_label.font = data_font_bold
sum_label.alignment = Alignment(horizontal="center", vertical="center")

for col_idx in range(1, 10):
    ws_skr1.cell(row=summary_row, column=col_idx).border = thin_border
    ws_skr1.cell(row=summary_row, column=col_idx).fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

b_cell = ws_skr1.cell(row=summary_row, column=3)
b_cell.value = total_budget
b_cell.font = data_font_bold
b_cell.number_format = money_format
b_cell.alignment = Alignment(horizontal="right", vertical="center")

a_cell = ws_skr1.cell(row=summary_row, column=7)
a_cell.value = f"รวมเงินตามสัญญา: {total_agreed:,.2f} บาท"
a_cell.font = data_font_bold
a_cell.alignment = Alignment(horizontal="center", vertical="center")

col_widths = {
    "A": 7,    # ลำดับ
    "B": 42,   # งานที่จัดซื้อ
    "C": 18,   # วงเงิน
    "D": 18,   # ราคากลาง
    "E": 20,   # วิธีซื้อจ้าง
    "F": 32,   # ผู้เสนอราคา
    "G": 32,   # ผู้ได้รับการคัดเลือก
    "H": 25,   # เหตุผล
    "I": 24    # เลขที่และวันที่สัญญา
}
for col_l, w in col_widths.items():
    ws_skr1.column_dimensions[col_l].width = w

# Save workbook
output_xlsx = "mockup-data-procurement.xlsx"
wb.save(output_xlsx)
print(f"Generated {output_xlsx} successfully!")

# Also generate CSV files for easy raw inspection / direct copy-paste
with open("mockup_contracts.csv", "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=contract_headers)
    writer.writeheader()
    writer.writerows(contracts)

with open("mockup_monthly_reports.csv", "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=report_headers)
    writer.writeheader()
    writer.writerows(reports)

with open("mockup_monthly_report_items.csv", "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=item_headers)
    writer.writeheader()
    writer.writerows(report_items)

print("Generated CSV files successfully!")
